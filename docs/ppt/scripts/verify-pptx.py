"""Verify an exported ppt-master PPTX: slide count, minimum rendered font size, notes.

Reads the raw DrawingML so the reported sizes are exactly what PowerPoint will use
(sz is in hundredths of a point: sz="1425" == 14.25 pt).
"""
from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

from pptx import Presentation


def main() -> int:
    pptx = Path(sys.argv[1])
    prs = Presentation(str(pptx))
    print(f"file    : {pptx.name}  ({pptx.stat().st_size / 1024:.0f} KB)")
    print(f"slides  : {len(prs.slides)}")

    sizes: list[tuple[float, str]] = []
    with zipfile.ZipFile(pptx) as zf:
        for name in zf.namelist():
            if not re.match(r"ppt/slides/slide\d+\.xml$", name):
                continue
            xml = zf.read(name).decode("utf-8")
            for m in re.finditer(r'sz="(\d+)"', xml):
                sizes.append((int(m.group(1)) / 100.0, name))
    sizes.sort()
    if sizes:
        print(f"min font: {sizes[0][0]:.2f} pt   (in {sizes[0][1]})")
        print(f"max font: {sizes[-1][0]:.2f} pt")
        below = [(s, n) for s, n in sizes if s < 14.0]
        print(f"runs below 14 pt: {len(below)}")
        small = sorted({s for s, _ in sizes})
        print("size ladder (pt): " + ", ".join(f"{s:g}" for s in small))

    missing = []
    for i, slide in enumerate(prs.slides, 1):
        text = ""
        if slide.has_notes_slide:
            text = slide.notes_slide.notes_text_frame.text.strip()
        if not text:
            missing.append(i)
    print(f"notes   : {len(prs.slides) - len(missing)}/{len(prs.slides)} slides"
          + (f"  missing: {missing}" if missing else ""))
    return 1 if missing or any(s < 14.0 for s, _ in sizes) else 0


if __name__ == "__main__":
    raise SystemExit(main())
